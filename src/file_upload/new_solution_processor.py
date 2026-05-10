import re
from typing import List, Dict, Optional, Tuple, Any
from datetime import datetime, date
import pandas as pd
import openpyxl
from sqlalchemy.orm import Session
from pathlib import Path

from src.database.models import NewSolution
from src.database.service import SQLiteService
from src.utils.logger import get_logger

logger = get_logger(__name__)


class NewSolutionProcessor:
    def __init__(self, db_service: SQLiteService):
        self.db_service = db_service

        self.column_mapping = {
            'urf_code': ['urf_code', 'урф_код', 'код урф', 'код_урф', 'урф код', 'уpf', 'код', 'urf', 'кодурф',
                         'урф-код'],
            'tb_name': ['тб', 'tb_name', 'территор', 'банк', 'tb', 'тербанк', 'территориальный', 'т/б'],
            'address': ['адрес', 'address', 'местоположение', 'адресс', 'локация', 'расположение'],
            'actual_decision_cs': ['решение по цс', 'actual_decision_cs', 'решение цс', 'цс решение',
                                   'актуальное решение', 'actual decision'],
            'actual_event_year': ['год мероприятия', 'actual_event_year', 'год', 'event_year', 'мероприятия год',
                                  'actual year'],
            'decision_kun_cs_vsp': ['решение кун', 'decision_kun_cs_vsp', 'кун решение', 'кун', 'kun',
                                    'решение кун цс'],
            'event_year_kun_cs_vsp': ['год мероприятия кун', 'event_year_kun_cs_vsp', 'кун год', 'кун мероприятия'],
            'closing_decision': ['решение о закрытии', 'closing_decision', 'закрытие решение', 'закрытия решение'],
            'closing_date': ['дата закрытия', 'closing_date', 'дата', 'closing', 'закрытия дата', 'date'],
            'comments': ['комментарии', 'comments', 'примечание', 'коммент', 'comment']
        }

    def process(self, file_path: Path) -> bool:
        try:
            if not isinstance(file_path, Path):
                file_path = Path(file_path)

            logger.info(f"Обработка файла: {file_path.name}")

            data = self._read_and_parse_file(file_path)
            if not data:
                logger.error("Не удалось прочитать данные из файла")
                return False

            with self.db_service.get_session() as session:
                self._save_to_database(session, data)

            logger.info(f"Успешно обработано {len(data)} записей")
            return True

        except Exception as e:
            logger.error(f"Ошибка при обработке файла: {str(e)}")
            import traceback
            logger.error(traceback.format_exc())
            return False

    def _read_and_parse_file(self, file_path: Path) -> List[Dict]:
        """Читает и парсит файл Excel"""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = self._find_data_sheet(wb)
            if not sheet:
                logger.error("Не найден лист с данными")
                return []

            logger.info(f"Используем лист: {sheet.title}")
            header_row, headers = self._find_headers(sheet)
            if header_row is None:
                logger.error("Не удалось найти заголовки")
                return []

            logger.info(f"Заголовки найдены в строке {header_row + 1}: {headers}")

            columns_map = self._build_columns_mapping(headers)
            logger.info(f"Маппинг колонок: {columns_map}")

            data = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row + 2, values_only=True),
                                          start=header_row + 2):

                if all(cell is None or (isinstance(cell, str) and cell.strip() == '') for cell in row):
                    continue

                record = {}
                for target_col, col_idx in columns_map.items():
                    if col_idx < len(row):
                        value = row[col_idx]
                        record[target_col] = self._clean_value(value, target_col)
                    else:
                        record[target_col] = None

                if record.get('urf_code') or record.get('address'):
                    data.append(record)

                if len(data) % 1000 == 0:
                    logger.info(f"Прочитано {len(data)} записей...")

            wb.close()
            logger.info(f"Всего прочитано {len(data)} записей")
            return data

        except Exception as e:
            logger.error(f"Ошибка чтения файла: {str(e)}")
            return []

    @staticmethod
    def _find_data_sheet(wb) -> Optional[Any]:
        """Находит лист с данными"""

        priority_sheets = ['решения', 'new', 'новые', 'кун', 'kun', 'solution', 'данные', 'data']

        for sheet_name in wb.sheetnames:
            lower_name = sheet_name.lower()
            for keyword in priority_sheets:
                if keyword in lower_name:
                    return wb[sheet_name]

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            data_count = 0
            for row in sheet.iter_rows(min_row=1, max_row=10, min_col=1, max_col=5, values_only=True):
                for cell in row:
                    if cell is not None and str(cell).strip():
                        data_count += 1
                        if data_count > 5:
                            return sheet
        return wb[wb.sheetnames[0]] if wb.sheetnames else None

    def _find_headers(self, sheet) -> Tuple[Optional[int], List[str]]:
        """Находит строку с заголовками"""
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            headers = []
            header_count = 0

            for cell in row:
                if cell is None:
                    headers.append('')
                    continue

                cell_str = str(cell).strip()
                headers.append(cell_str)

                if cell_str:
                    if 2 <= len(cell_str) <= 50:
                        cell_lower = cell_str.lower()
                        for aliases in self.column_mapping.values():
                            for alias in aliases:
                                if alias in cell_lower or cell_lower in alias:
                                    header_count += 1
                                    break

            if header_count >= 3:
                logger.info(f"Найдена строка заголовков: строка {row_idx}")
                return row_idx - 1, headers

        logger.info("Заголовки не найдены, используем первую строку")
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        return 0, [str(cell).strip() if cell else '' for cell in first_row]

    def _build_columns_mapping(self, headers: List[str]) -> Dict[str, int]:
        """Строит маппинг колонок по заголовкам"""
        mapping = {}
        used_columns = set()

        for target_col, aliases in self.column_mapping.items():
            best_match_idx = -1
            best_match_score = 0

            for col_idx, header in enumerate(headers):
                if not header or col_idx in used_columns:
                    continue

                header_lower = header.lower().strip()

                for alias in aliases:
                    alias_lower = alias.lower()

                    if alias_lower == header_lower:
                        score = 100
                    elif alias_lower in header_lower:
                        score = 80
                    elif header_lower in alias_lower:
                        score = 60
                    elif self._similar(alias_lower, header_lower):
                        score = 40
                    else:
                        continue

                    if score > best_match_score:
                        best_match_score = score
                        best_match_idx = col_idx

            if best_match_idx >= 0 and best_match_score > 30:
                mapping[target_col] = best_match_idx
                used_columns.add(best_match_idx)
                logger.debug(f"Сопоставлено: '{headers[best_match_idx]}' -> {target_col}")

        if 'urf_code' not in mapping and headers:
            first_header = headers[0].lower() if headers[0] else ''
            if any(keyword in first_header for keyword in ['код', 'урф', 'code', 'id']):
                mapping['urf_code'] = 0
                logger.info("URF_CODE назначен на первую колонку")

        return mapping

    @staticmethod
    def _similar(str1: str, str2: str) -> bool:
        """Проверяет, похожи ли строки"""
        str1 = re.sub(r'[_\s\-]', '', str1.lower())
        str2 = re.sub(r'[_\s\-]', '', str2.lower())
        return str1 in str2 or str2 in str1

    def _clean_value(self, value, field_name: str) -> Any:
        """Очищает и преобразует значение"""
        if value is None:
            return None

        if isinstance(value, str):
            value = value.strip()
            if not value or value.lower() in ['', 'nan', 'none', 'null']:
                return None

        if field_name == 'closing_date':
            return self._parse_date(value)
        elif field_name in ['actual_event_year', 'event_year_kun_cs_vsp']:
            return self._parse_year(value)

        if isinstance(value, (datetime, pd.Timestamp)):
            return str(value.date()) if field_name != 'closing_date' else value.date()
        elif isinstance(value, (int, float)):
            return str(int(value)) if field_name in ['actual_event_year', 'event_year_kun_cs_vsp'] else str(value)

        return str(value) if value else None

    @staticmethod
    def _parse_date(value: Any) -> Optional[date]:
        """Парсит дату и возвращает объект date"""
        if not value:
            return None

        if isinstance(value, (datetime, pd.Timestamp)):
            return value.date()

        if isinstance(value, date):
            return value

        if isinstance(value, str):
            value = value.strip()
            if not value:
                return None

            formats = [
                '%Y-%m-%d',
                '%d.%m.%Y', '%d/%m/%Y', '%d-%m-%Y',
                '%Y.%m.%d', '%Y/%m/%d',
                '%d.%m.%y', '%d/%m/%y', '%d-%m-%y'
            ]

            for fmt in formats:
                try:
                    dt = datetime.strptime(value, fmt)
                    if fmt.endswith('%y') and dt.year < 2000:
                        dt = dt.replace(year=dt.year + 2000)
                    return dt.date()
                except ValueError:
                    continue

            if '-' in value or '.' in value or '/' in value:
                parts = re.findall(r'\d+', value)
                if len(parts) >= 3:
                    try:
                        if len(parts[0]) == 4:
                            year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
                        else:
                            day, month, year = int(parts[0]), int(parts[1]), int(parts[2])

                        if year < 100:
                            year += 2000 if year < 50 else 1900

                        return date(year, month, day)
                    except (ValueError, IndexError):
                        pass

        return None

    @staticmethod
    def _parse_year(value: Any) -> Optional[str]:
        """Парсит год, возвращает строку"""
        if not value:
            return None

        if isinstance(value, (int, float)):
            year = int(value)
            if 1900 <= year <= 2100:
                return str(year)

        if isinstance(value, str):
            value = value.strip()
            match = re.search(r'(19\d{2}|20\d{2})', value)
            if match:
                return match.group(1)

            if value.isdigit() and len(value) == 4:
                year = int(value)
                if 1900 <= year <= 2100:
                    return value

        return str(value) if value else None

    @staticmethod
    def _save_to_database(session: Session, data: List[Dict[str, Any]]) -> None:
        if not data:
            logger.info("Нет данных для сохранения")
            return

        try:
            deleted_count = session.query(NewSolution).delete()
            session.flush()
            logger.info(f"Удалено {deleted_count} старых записей")

            records_to_insert = []
            for item in data:
                record = {
                    'urf_code': item.get('urf_code'),
                    'tb_name': item.get('tb_name'),
                    'address': item.get('address'),
                    'actual_decision_cs': item.get('actual_decision_cs'),
                    'actual_event_year': item.get('actual_event_year'),
                    'decision_kun_cs_vsp': item.get('decision_kun_cs_vsp'),
                    'event_year_kun_cs_vsp': item.get('event_year_kun_cs_vsp'),
                    'closing_decision': item.get('closing_decision'),
                    'closing_date': item.get('closing_date'),
                    'comments': item.get('comments')
                }
                records_to_insert.append(record)

            if records_to_insert:
                session.bulk_insert_mappings(NewSolution, records_to_insert)
                session.commit()
                logger.info(f"Сохранено {len(records_to_insert)} новых записей")

        except Exception as e:
            session.rollback()
            logger.error(f"Ошибка сохранения в базу: {str(e)}")
            raise


def process_new_solution_file(file_path: str) -> bool:
    """Основная функция обработки файла"""
    logger.info("Запуск обработки файла с новыми решениями")

    file_path = Path(file_path)
    if not file_path.exists():
        logger.error(f"Файл не найден: {file_path}")
        return False

    logger.info(f"Файл: {file_path.name}")

    try:
        db_service = SQLiteService()
        processor = NewSolutionProcessor(db_service)

        success = processor.process(file_path)

        if success:
            logger.info("Обработка завершена успешно!")
        else:
            logger.error("Обработка завершена с ошибками!")

        return success

    except Exception as e:
        logger.error(f"Критическая ошибка: {str(e)}")
        import traceback
        logger.error(traceback.format_exc())
        return False


if __name__ == "__main__":
    import sys

    project_root = Path(__file__).resolve().parents[2]
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = str(project_root / "data" / "Новые решения.xlsx")

    print(f"📂 Обработка файла: {file_path}")
    success = process_new_solution_file(file_path)

    if success:
        print("✅ Готово!")
        sys.exit(0)
    else:
        print("❌ Ошибка!")
        sys.exit(1)