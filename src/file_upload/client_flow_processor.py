import re
from typing import List, Dict, Optional, Any
from datetime import datetime, date
import openpyxl
from pathlib import Path

from src.database.models import ClientFlow
from src.database.service import SQLiteService
from src.utils.logger import get_logger

logger = get_logger(__name__)


class ClientFlowProcessor:
    def __init__(self, db_service: SQLiteService):
        self.db_service = db_service
        self.column_mapping = {
            'urf_code': ['urf_code', 'урф код', 'код урф', 'urf', 'код'],
            'tb_code': ['# тб', '#тб', 'код тб'],  # 13
            'tb': ['тб', 'tb', 'территориальный банк', 'тербанк'],  # ЦЧБ
            'gosb': ['госб', 'gosb'],
            'city_id': ['city_id', 'id города', 'код города'],
            'city': ['город', 'city', 'населенный пункт'],
            'type_np': ['тип нп', 'type_np', 'тип'],
            'year': ['год', 'year', 'г'],
            'month': ['месяц', 'month', 'мес', 'дата'],  # ИСПРАВЛЕНО: month вместо month_date
            'up': ['уп', 'up'],
            'kp_skm': ['кп скм', 'kp_skm', 'скм'],
            'kp_smo': ['кп смо', 'kp_smo', 'смо'],
            'total_kp': ['общий кп', 'total_kp', 'итого кп', 'всего кп']
        }

    def process(self, file_path: Path) -> bool:
        try:
            file_path = Path(file_path)
            logger.info(f"Обработка файла: {file_path.name}")

            data = self._read_excel(file_path)
            if not data:
                logger.error("Нет данных для обработки")
                return False

            # Отладка: выведем первые 2 записи чтобы увидеть даты
            print("\n=== ПЕРВЫЕ 2 ЗАПИСИ ИЗ EXCEL ===")
            for i, record in enumerate(data[:2]):
                print(f"Запись {i + 1}:")
                print(f"  urf_code: {record.get('urf_code')}")
                print(f"  month: {record.get('month')} (тип: {type(record.get('month'))})")
                print(f"  year: {record.get('year')}")
            print("================================\n")

            with self.db_service.get_session() as session:
                # Очистка таблицы
                deleted_count = session.query(ClientFlow).delete()
                print(f"Удалено {deleted_count} записей из таблицы ClientFlow")
                session.flush()

                # Вставка новых данных
                if data:
                    session.bulk_insert_mappings(ClientFlow, data)
                    session.commit()
                    logger.info(f"Сохранено {len(data)} записей")

            logger.info(f"Успешно обработано {len(data)} записей")
            return True

        except Exception as e:
            logger.error(f"Ошибка: {str(e)}")
            return False

    def _read_excel(self, file_path: Path) -> List[Dict]:
        """Чтение Excel файла"""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = wb.active

            # Поиск заголовков (первые 10 строк)
            header_row = None
            headers = []
            for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), 1):
                row_values = [str(cell).strip() if cell else '' for cell in row]
                if any(self._is_header_cell(val) for val in row_values[:3]):
                    header_row = i
                    headers = row_values
                    break

            if not header_row:
                header_row = 1
                headers = [str(cell).strip() if cell else '' for cell in next(sheet.iter_rows(values_only=True))]

            logger.info(f"Заголовки: {headers}")

            # Маппинг колонок
            col_map = {}
            for col_idx, header in enumerate(headers):
                header_lower = header.lower()
                for field, aliases in self.column_mapping.items():
                    if any(alias.lower() in header_lower for alias in aliases):
                        col_map[field] = col_idx
                        print(f"Найдено соответствие: '{header}' -> {field}")
                        break

            logger.info(f"Маппинг: {col_map}")

            # Чтение данных
            data = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True)):
                if all(cell is None for cell in row):
                    continue

                record = {}
                for field, col_idx in col_map.items():
                    if col_idx < len(row):
                        value = row[col_idx]
                        original_value = value
                        cleaned_value = self._clean_value(value, field)
                        record[field] = cleaned_value

                        # Отладка для поля month
                        if field == 'month' and row_idx < 2:  # только первые 2 строки
                            print(f"Строка {row_idx + header_row + 1}, поле {field}:")
                            print(f"  исходное: {original_value} (тип: {type(original_value)})")
                            print(f"  очищенное: {cleaned_value} (тип: {type(cleaned_value)})")

                if record.get('urf_code'):
                    data.append(record)

            wb.close()
            return data

        except Exception as e:
            logger.error(f"Ошибка чтения Excel: {str(e)}")
            return []

    def _is_header_cell(self, value: Any) -> bool:
        """Проверка, является ли ячейка заголовком"""
        if not value or not isinstance(value, str):
            return False
        value_lower = value.lower()
        keywords = ['urf', 'код', 'тб', 'город', 'год', 'месяц', 'уп', 'кп']
        return any(kw in value_lower for kw in keywords)

    def _clean_value(self, value: Any, field: str) -> Any:
        """Очистка значения"""
        if value is None:
            return None

        if isinstance(value, str):
            value = value.strip()
            if not value or value.lower() in ['', 'nan', 'null', 'none']:
                return None

        # Для поля с датой
        if field == 'month':
            # Пробуем разные форматы
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            if isinstance(value, str):
                # Пробуем основные форматы
                for fmt in ['%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d', '%d.%m.%y']:
                    try:
                        dt = datetime.strptime(value, fmt)
                        return dt.date()
                    except ValueError:
                        continue
            return None

        # Для года
        if field == 'year':
            year_val = self._extract_year(value)
            return str(year_val) if year_val else None

        # Для числовых полей (коды, показатели)
        if field in ['tb_code', 'gosb', 'city_id', 'up', 'kp_skm', 'kp_smo', 'total_kp']:
            return self._parse_number(value)

        # Для текстовых полей
        return str(value) if value else None

    def _extract_year(self, value: Any) -> Optional[int]:
        """Извлекает год"""
        if not value:
            return None

        if isinstance(value, (int, float)):
            return int(value)

        if isinstance(value, str):
            # Ищем 4 цифры подряд
            match = re.search(r'20\d{2}|19\d{2}', value)
            if match:
                return int(match.group())

        return None

    def _parse_number(self, value: Any) -> Optional[int]:
        """Парсинг числа"""
        if value is None:
            return None

        if isinstance(value, (int, float)):
            return int(value)

        if isinstance(value, str):
            # Убираем все кроме цифр и минуса
            value = re.sub(r'[^\d\-]', '', value)
            if value and value.lstrip('-').isdigit():
                return int(value)

        return None


def process_client_flow_file(file_path: str) -> bool:
    """Основная функция обработки"""
    logger.info("Запуск обработки файла с клиентским потоком")

    file_path = Path(file_path)
    if not file_path.exists():
        logger.error(f"Файл не найден: {file_path}")
        return False

    try:
        db_service = SQLiteService()
        processor = ClientFlowProcessor(db_service)
        return processor.process(file_path)

    except Exception as e:
        logger.error(f"Ошибка: {str(e)}")
        return False


if __name__ == "__main__":
    import sys

    project_root = Path(__file__).resolve().parents[2]
    default_path = str(project_root / "data" / "client_flow.xlsx")
    file_path = sys.argv[1] if len(sys.argv) > 1 else default_path
    success = process_client_flow_file(file_path)
    print("✅ Готово!" if success else "❌ Ошибка!")
    sys.exit(0 if success else 1)
